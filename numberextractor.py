import sys
import os
import re
import docx
import openpyxl
import pandas as pd
from PyQt6.QtWidgets import QApplication, QMainWindow, QFileDialog, QTextEdit, QPushButton, QMessageBox

class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()

        # Initialize file_name as an instance variable
        self.file_name = ""

        self.init_ui()

    def init_ui(self):
        self.text_edit = QTextEdit(self)
        self.setCentralWidget(self.text_edit)

        open_file_button = QPushButton("Open File", self)
        open_file_button.clicked.connect(self.open_file)
        open_file_button.setGeometry(50, 50, 80, 30)

        extract_button = QPushButton("Extract Numbers", self)
        extract_button.clicked.connect(self.extract_numbers)
        extract_button.setGeometry(150, 50, 100, 30)

    def open_file(self):
        # Use self.file_name to store the selected file path
        self.file_name, _ = QFileDialog.getOpenFileName(self, "Open File", "", "Text Files (*.txt);;Excel Files (*.xlsx);;Word Files (*.docx)")

        if self.file_name:
            with open(self.file_name, "r") as f:
                self.text_edit.setText(f.read())

    def extract_numbers(self):
        text = self.text_edit.toPlainText()

        # Extract numbers from text
        numbers = re.findall(r"\d+", text)

        # Extract numbers from Excel file
        if ".xlsx" in self.file_name:
            wb = openpyxl.load_workbook(self.file_name)
            sheet = wb.active
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value and isinstance(cell.value, (int, float)):
                        numbers.append(str(cell.value))

        # Extract numbers from Word file
        elif ".docx" in self.file_name:
            doc = docx.Document(self.file_name)
            for para in doc.paragraphs:
                for run in para.runs:
                    if run.text and re.match(r"\d+", run.text):
                        numbers.append(run.text)

        # Extract numbers from text file
        elif ".txt" in self.file_name:
            with open(self.file_name, "r") as f:
                for line in f:
                    for word in line.split():
                        if word.isdigit():
                            numbers.append(word)

        # Display extracted numbers
        if numbers:
            self.text_edit.setText("\n".join(numbers))
        else:
            QMessageBox.warning(self, "Warning", "No numbers found in file.")

if __name__ == "__main__":
    app = QApplication(sys.argv)
    window = MainWindow()
    window.show()
    sys.exit(app.exec())
